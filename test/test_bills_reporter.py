"""Unit tests for cogs/bills/reporter.py"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from cogs.bills.reporter import build_all_report, build_single_report

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

_OLD_BILL = {
    "bill_number": "S1234",
    "state": "NJ",
    "title": "A Bill Concerning Testing",
    "description": "This bill is about automated testing.",
    "status": "Introduced",
    "last_action": "Introduced in the Senate",
    "last_action_date": "2026-01-15",
    "last_alerted_action_date": "2026-01-15",  # same → NOT new
    "change_hash": "abc",
    "sponsors": ["Sen. Smith", "Sen. Jones"],
    "url": "https://www.njleg.state.nj.us/bill-search/2026/S1234",
    "added_by": "testuser",
    "added_at": "2026-01-10T12:00:00",
    "last_checked": "2026-01-15T08:00:00",
    "history": [
        {"date": "2026-01-10", "action": "Introduced", "chamber": "S", "importance": 1},
        {"date": "2026-01-15", "action": "Referred to Committee", "chamber": "S", "importance": 0},
    ],
}

_NEW_BILL = {
    **_OLD_BILL,
    "bill_number": "A567",
    "title": "Another New Bill",
    "last_action": "Passed Assembly",
    "last_action_date": "2026-03-20",
    "last_alerted_action_date": "2026-01-15",  # older → IS new
    "history": [
        {"date": "2026-01-15", "action": "Introduced", "chamber": "A", "importance": 1},
        {"date": "2026-03-20", "action": "Passed Assembly", "chamber": "A", "importance": 1},
    ],
}


def _load(xlsx_bytes: bytes):
    return load_workbook(io.BytesIO(xlsx_bytes))


def _is_yellow(cell) -> bool:
    fill = cell.fill
    # After xlsx round-trip openpyxl wraps fills in StyleProxy (not PatternFill),
    # so check fill_type and color directly without isinstance.
    return fill.fill_type == "solid" and fill.fgColor.rgb in ("FFFF00", "00FFFF00")


# ---------------------------------------------------------------------------
# Single-bill report tests
# ---------------------------------------------------------------------------

def test_single_report_returns_bytes():
    result = build_single_report(_OLD_BILL)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_single_report_sheets():
    wb = _load(build_single_report(_OLD_BILL))
    assert "Overview" in wb.sheetnames
    assert "History" in wb.sheetnames


def test_single_report_overview_fields():
    wb = _load(build_single_report(_OLD_BILL))
    ws = wb["Overview"]
    # Collect all cell values
    values = {str(cell.value) for row in ws.iter_rows() for cell in row if cell.value}
    assert "S1234" in values
    assert "A Bill Concerning Testing" in values
    assert "Introduced" in values


def test_single_report_history_new_rows_highlighted():
    """Rows in History newer than last_alerted_action_date should be yellow."""
    bill = {
        **_OLD_BILL,
        "last_alerted_action_date": "2026-01-12",  # only 2026-01-15 row is newer
        "history": [
            {"date": "2026-01-10", "action": "Introduced", "chamber": "S", "importance": 1},
            {"date": "2026-01-15", "action": "Referred to Committee", "chamber": "S", "importance": 0},
        ],
    }
    wb = _load(build_single_report(bill))
    ws = wb["History"]

    # Row 2 (date 2026-01-10) should NOT be yellow
    assert not _is_yellow(ws.cell(row=2, column=1))
    # Row 3 (date 2026-01-15) should be yellow
    assert _is_yellow(ws.cell(row=3, column=1))


def test_single_report_history_old_rows_not_highlighted():
    """Rows not newer than last_alerted_action_date should have no yellow fill."""
    wb = _load(build_single_report(_OLD_BILL))
    ws = wb["History"]
    # All history dates == last_alerted_action_date or older → none highlighted
    for row in range(2, ws.max_row + 1):
        assert not _is_yellow(ws.cell(row=row, column=1))


# ---------------------------------------------------------------------------
# All-bills report tests
# ---------------------------------------------------------------------------

def test_all_report_returns_bytes():
    result = build_all_report([_OLD_BILL, _NEW_BILL])
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_all_report_sheet():
    wb = _load(build_all_report([_OLD_BILL, _NEW_BILL]))
    assert "Bills" in wb.sheetnames


def test_all_report_changed_row_highlighted():
    """Bill with last_action_date > last_alerted_action_date should be yellow."""
    wb = _load(build_all_report([_OLD_BILL, _NEW_BILL]))
    ws = wb["Bills"]

    # Row 2 = _OLD_BILL (no change), Row 3 = _NEW_BILL (changed)
    assert not _is_yellow(ws.cell(row=2, column=1))
    assert _is_yellow(ws.cell(row=3, column=1))


def test_all_report_unchanged_row_not_highlighted():
    """Bill with last_action_date == last_alerted_action_date should not be yellow."""
    wb = _load(build_all_report([_OLD_BILL]))
    ws = wb["Bills"]
    assert not _is_yellow(ws.cell(row=2, column=1))


def test_all_report_single_bill():
    """build_all_report works with a single-element list."""
    result = build_all_report([_OLD_BILL])
    wb = _load(result)
    ws = wb["Bills"]
    assert ws.max_row == 2  # header + 1 data row
