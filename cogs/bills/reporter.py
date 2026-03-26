"""
Excel report generation — pure openpyxl, no network or Discord imports.

To add new report types, add functions here.  cog.py calls them and
uploads the returned bytes as a Discord file attachment.

Public API:
  build_single_report(bill: dict) -> bytes
  build_all_report(bills: list[dict]) -> bytes
"""

from __future__ import annotations

import io
import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Yellow highlight for changed / new rows
_YELLOW = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")


# ---------------------------------------------------------------------------
# Public functions
# ---------------------------------------------------------------------------

def build_single_report(bill: dict[str, Any]) -> bytes:
    """
    Build a single-bill Excel report.

    Sheets:
      Overview  — key fields (number, title, sponsors, status, description, URL)
      History   — table of all legislative actions; rows newer than
                  last_alerted_action_date are highlighted yellow
    """
    wb = Workbook()

    _build_overview_sheet(wb.active, bill)
    wb.active.title = "Overview"

    hist_ws = wb.create_sheet("History")
    _build_history_sheet(hist_ws, bill)

    return _to_bytes(wb)


def build_all_report(bills: list[dict[str, Any]]) -> bytes:
    """
    Build a multi-bill summary Excel report.

    Single sheet 'Bills':
      Bill # | Title | Status | Last Action | Last Action Date | Recent History (last 3) | Sponsors
    Rows where last_action_date > last_alerted_action_date are highlighted yellow.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Bills"

    headers = ["Bill #", "Title", "Status", "Last Action", "Last Action Date", "Recent History (last 3)", "Sponsors"]
    _write_header_row(ws, headers)

    # Column index for "Recent History (last 3)" — used to enable wrap_text
    _HISTORY_COL = 6

    for bill in bills:
        recent_history = bill.get("history", [])[-3:]
        history_text = "\n".join(
            f"{h.get('date', '')} — {h.get('action', '')}"
            for h in reversed(recent_history)
        )
        row = [
            bill.get("bill_number", ""),
            bill.get("title", ""),
            bill.get("status", ""),
            bill.get("last_action", ""),
            bill.get("last_action_date", ""),
            history_text,
            ", ".join(bill.get("sponsors", [])),
        ]
        ws.append(row)
        ws.cell(row=ws.max_row, column=_HISTORY_COL).alignment = Alignment(wrap_text=True)

        # Highlight row if there's a new action since the last alert
        last_date = bill.get("last_action_date", "")
        alerted_date = bill.get("last_alerted_action_date", "")
        if last_date and last_date > alerted_date:
            _highlight_row(ws, ws.max_row, len(headers))

    _auto_width(ws)
    ws.freeze_panes = "A2"

    return _to_bytes(wb)


# ---------------------------------------------------------------------------
# Private sheet builders
# ---------------------------------------------------------------------------

def _build_overview_sheet(ws, bill: dict) -> None:
    fields = [
        ("Bill Number", bill.get("bill_number", "")),
        ("State", bill.get("state", "")),
        ("Title", bill.get("title", "")),
        ("Status", bill.get("status", "")),
        ("Last Action", bill.get("last_action", "")),
        ("Last Action Date", bill.get("last_action_date", "")),
        ("Sponsors", ", ".join(bill.get("sponsors", []))),
        ("Description", bill.get("description", "")),
        ("URL", bill.get("url", "")),
        ("Added By", bill.get("added_by", "")),
        ("Added At", bill.get("added_at", "")),
        ("Last Checked", bill.get("last_checked", "")),
    ]

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 80

    for label, value in fields:
        ws.append([label, str(value) if value else ""])
        label_cell = ws.cell(row=ws.max_row, column=1)
        label_cell.font = _HEADER_FONT
        value_cell = ws.cell(row=ws.max_row, column=2)
        value_cell.alignment = Alignment(wrap_text=True)


def _build_history_sheet(ws, bill: dict) -> None:
    headers = ["Date", "Chamber", "Action", "Importance"]
    _write_header_row(ws, headers)

    alerted_date = bill.get("last_alerted_action_date", "")

    for entry in bill.get("history", []):
        row = [
            entry.get("date", ""),
            entry.get("chamber", ""),
            entry.get("action", ""),
            entry.get("importance", 0),
        ]
        ws.append(row)
        # Highlight rows that are newer than the last alerted date
        if entry.get("date", "") > alerted_date:
            _highlight_row(ws, ws.max_row, len(headers))

    _auto_width(ws)
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _write_header_row(ws, headers: list[str]) -> None:
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def _highlight_row(ws, row_num: int, num_cols: int) -> None:
    for col in range(1, num_cols + 1):
        ws.cell(row=row_num, column=col).fill = _YELLOW


def _auto_width(ws) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
